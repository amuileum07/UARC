"""Bounded numeric CSV/XLSX profiles; no workbook execution or arbitrary paths."""

import csv
import hashlib
import io
import re
from datetime import UTC, datetime
from itertools import islice
from math import isfinite
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from rocketpy import Environment
from rocketpy.tools import geopotential_height_to_geometric_height

from app.schemas import WeatherImportOptions, WeatherProfile, WeatherTable
from app.weather import wind_components

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_ROWS = 4096
MAX_COLUMNS = 32


def _xlsx_rows(raw, sheet):
    from defusedxml.ElementTree import fromstring
    from openpyxl import DEFUSEDXML, load_workbook

    if not DEFUSEDXML:
        raise RuntimeError("The XLSX reader requires defusedxml.")
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 128
                or sum(item.file_size for item in entries) > 16 * 1024 * 1024
            ):
                raise ValueError("Excel 압축 해제 크기·항목 제한을 초과했습니다.")
            if any(
                item.flag_bits & 1 or "vbaproject" in item.filename.lower()
                for item in entries
            ):
                raise ValueError("암호화·매크로 통합 문서는 지원하지 않습니다.")
            # Reject hidden out-of-range cells before a bounded reader could drop
            # them. Export only the profile sheet if a workbook exceeds these limits.
            for item in entries:
                if item.filename.startswith(
                    "xl/worksheets/"
                ) and item.filename.endswith(".xml"):
                    root = fromstring(archive.read(item), forbid_dtd=True)
                    for cell in root.iter(
                        "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"
                    ):
                        match = re.fullmatch(
                            r"([A-Z]+)([1-9][0-9]*)", cell.get("r", "")
                        )
                        if match is None:
                            raise ValueError("Invalid Excel cell reference.")
                        letters, row = match.groups()
                        column = 0
                        for letter in letters:
                            column = column * 26 + ord(letter) - ord("A") + 1
                        if column > MAX_COLUMNS or int(row) > MAX_ROWS + 1:
                            raise ValueError(
                                "Excel 데이터 범위는 열 32개·데이터 4096행까지입니다."
                            )
        workbook = load_workbook(
            io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
        )
    except BadZipFile as error:
        raise ValueError("올바른 .xlsx 파일이 아닙니다.") from error
    try:
        sheets = workbook.sheetnames
        if not sheets or len(sheets) > 16:
            raise ValueError("Excel 시트는 1~16개여야 합니다.")
        selected = sheet or sheets[0]
        if selected not in sheets:
            raise ValueError("선택한 시트가 없습니다.")
        worksheet = workbook[selected]
        worksheet.reset_dimensions()
        rows = []
        for cells in islice(worksheet.iter_rows(max_col=MAX_COLUMNS + 1), MAX_ROWS + 2):
            if any(cell.data_type in {"f", "e"} for cell in cells):
                raise ValueError(
                    "수식·오류 셀은 가져올 수 없습니다. 값으로 붙여넣은 파일을 사용하세요."
                )
            rows.append([cell.value for cell in cells])
        return rows, sheets, selected
    finally:
        workbook.close()


def _read_table(source: Path, sheet: str | None = None):
    with source.open("rb") as stream:
        raw = stream.read(MAX_FILE_BYTES + 1)
    if len(raw) > MAX_FILE_BYTES:
        raise ValueError("기상 파일은 최대 8 MiB입니다.")
    suffix = source.suffix.lower()
    if suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp949")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(islice(csv.reader(io.StringIO(text), dialect), MAX_ROWS + 2))
        if len(rows) > MAX_ROWS + 1:
            raise ValueError("CSV 데이터는 최대 4096행입니다.")
        sheets, selected = [], None
    elif suffix == ".xlsx":
        rows, sheets, selected = _xlsx_rows(raw, sheet)
    else:
        raise ValueError(
            ".csv 또는 .xlsx를 선택하세요. .xls는 .xlsx로 저장한 뒤 가져오세요."
        )
    if not rows:
        raise ValueError("빈 기상 파일입니다.")

    def trimmed(row):
        row = list(row)
        while row and (row[-1] is None or str(row[-1]).strip() == ""):
            row.pop()
        return row

    columns = [
        str(cell).strip() if cell is not None else "" for cell in trimmed(rows[0])
    ]
    if (
        not 5 <= len(columns) <= MAX_COLUMNS
        or any(not c or len(c) > 128 for c in columns)
        or len(set(columns)) != len(columns)
    ):
        raise ValueError("첫 행에 중복·빈 이름 없이 5~32개 열 이름을 입력하세요.")
    records = []
    for row_number, row in enumerate(rows[1:], 2):
        row = trimmed(row)
        if not row:
            continue
        if row_number > MAX_ROWS + 1 or len(row) > len(columns):
            raise ValueError(
                "기상 파일은 최대 4096행이며 모든 값에 열 이름이 필요합니다."
            )
        records.append((row_number, row + [None] * (len(columns) - len(row))))
    table = WeatherTable(
        sheets=sheets,
        sheet=selected,
        columns=columns,
        rows=[
            [None if value is None else str(value)[:256] for value in row]
            for _, row in records[:5]
        ],
        row_count=len(records),
        sha256=hashlib.sha256(raw).hexdigest(),
    )
    return table, records


def preview_weather_file(source: Path, sheet: str | None = None) -> WeatherTable:
    return _read_table(source, sheet)[0]


def import_weather_file(source: Path, options: WeatherImportOptions) -> WeatherProfile:
    table, records = _read_table(source, options.sheet)
    if table.sha256 != options.expected_sha256:
        raise ValueError("미리보기 이후 파일이 바뀌었습니다. 파일을 다시 선택하세요.")
    columns = options.columns
    names = {
        "height": columns.altitude,
        "pressure": columns.pressure,
        "temperature": columns.temperature,
    }
    names |= (
        {"u": columns.wind_u, "v": columns.wind_v}
        if options.wind_format == "uv"
        else {"speed": columns.wind_speed, "direction": columns.wind_direction}
    )
    if any(name not in table.columns for name in names.values()):
        raise ValueError("선택한 열을 파일에서 찾을 수 없습니다.")
    indices = {key: table.columns.index(name) for key, name in names.items()}
    request = options.request
    radius = float(Environment.calculate_earth_radius(request.latitude))
    offset = (
        request.geoid_undulation_m if request.elevation_reference == "ellipsoid" else 0
    )
    wind_factor = {"m/s": 1, "km/h": 1 / 3.6, "knots": 1852 / 3600}[options.wind_unit]
    levels = []
    for row_number, row in records:
        values = {}
        for key, index in indices.items():
            value = row[index]
            try:
                if value is None or isinstance(value, bool):
                    raise ValueError("missing/boolean")
                number = float(value)
                if not isfinite(number):
                    raise ValueError("non-finite")
            except (ValueError, TypeError) as error:
                raise ValueError(
                    f"{row_number}행 {table.columns[index]}: 유효한 숫자가 필요합니다."
                ) from error
            values[key] = number
        height = values["height"] * (0.3048 if options.height_unit == "ft" else 1)
        if options.height_reference == "AGL":
            height += request.elevation - request.pad_height_m
        else:
            if options.height_reference == "geopotential":
                height = float(geopotential_height_to_geometric_height(height, radius))
            height += offset
        u, v = (
            (values["u"] * wind_factor, values["v"] * wind_factor)
            if options.wind_format == "uv"
            else wind_components(values["speed"] * wind_factor, values["direction"])
        )
        levels.append(
            dict(
                altitude_m=height,
                pressure_pa=values["pressure"]
                * (100 if options.pressure_unit == "hPa" else 1),
                temperature_k=values["temperature"]
                + (273.15 if options.temperature_unit == "C" else 0),
                wind_u_m_s=u,
                wind_v_m_s=v,
            )
        )
    levels.sort(key=lambda level: level["altitude_m"])
    return WeatherProfile(
        request=request,
        source="File",
        source_name=options.source_name,
        source_asset=str(source),
        retrieved_at=datetime.now(UTC),
        grid_latitude=request.latitude,
        grid_longitude=request.longitude,
        response_sha256=table.sha256,
        levels=levels,
        notes=[
            "파일의 고도 순서를 정렬했습니다. 중복 고도·기압 역전·결측값은 허용하지 않습니다.",
            "자료의 적용 위치·시각을 현재 발사 설정에 연결했습니다. 실제 관측 위치·시각은 파일 출처와 대조하세요.",
        ],
    )
